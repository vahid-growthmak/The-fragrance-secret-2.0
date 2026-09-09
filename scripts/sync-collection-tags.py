#!/usr/bin/env python3
"""Sync product membership of the scent/lifestyle/climate collections from the
listings spreadsheet.

Every one of these collections is a *smart* collection whose rule is a single
tag equality (verified against the store, not assumed). So there is no such
thing as "adding a product to the collection" through the API: membership is
decided by whether the product carries the tag. This script therefore writes
product tags, and nothing else.

    Oud & Woody          tag equals family-oud-woody
    Office Wear          tag equals office-wear
    Crazy Deals          tag equals sale
    ...

Per docs/GO-LIVE.md, a product-touching script must name the exact fields it
writes and touch nothing else. This one writes `tags` only, and within tags it
only ever adds or removes the 29 managed tags below — any other tag a product
carries is preserved verbatim.

Usage
-----
    python3 scripts/sync-collection-tags.py                 # dry run (default)
    python3 scripts/sync-collection-tags.py --report out.md # dry run + written report
    python3 scripts/sync-collection-tags.py --collection wedding   # one collection
    python3 scripts/sync-collection-tags.py --apply         # write (asks to confirm)

Credentials come from .env in the repo root, same as the other scripts:
    SHOPIFY_STORE_DOMAIN, SHOPIFY_ADMIN_ACCESS_TOKEN, SHOPIFY_API_VERSION

Nothing is written without --apply.
"""

import argparse
import json
import os
import re
import ssl
import sys
import time
import urllib.error
import urllib.request
from collections import defaultdict
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
XLSX = REPO / "Fragrance_Secrets_Product Listings.xlsx"

# sheet name -> (collection handle, the tag its smart rule matches)
# Beach/Vacation and Evening/Formal are easy to transpose, so these were taken
# from the live smart-collection rules rather than inferred from the names.
SHEETS = {
    "Oud Woody":        ("oud-woody",          "family-oud-woody"),
    "Fresh Citrus":     ("fresh-citrus",       "family-fresh-citrus"),
    "Floral Rose":      ("floral-rose",        "family-floral-rose"),
    "Sweet Gourmand":   ("sweet-gourmand",     "family-sweet-gourmand"),
    "Spicy Oriental":   ("spicy-oriental",     "family-spicy-oriental"),
    "Everyday":         ("everyday-signature", "everyday-signature"),
    "Office Wear":      ("office-wear",        "office-wear"),
    "Date Night":       ("date-night",         "date-night"),
    "Party Wear":       ("party-clubbing",     "party-clubbing"),
    "Wedding":          ("wedding",            "wedding"),
    "Formal Event":     ("wedding-formal",     "wedding-formal"),
    "Gym":              ("gym",                "gym"),
    "Travel":           ("travel-friendly",    "travel-friendly"),
    "Beach":            ("beach-vacation",     "beach-vacation"),
    "Vacation":         ("vacation-beach",     "vacation-beach"),
    "Evening Wear":     ("dinner-evening",     "dinner-evening"),
    "Climate-Summer":   ("summer-wear",        "summer-wear"),
    "Climate-Winter":   ("winter-holiday",     "winter-holiday"),
    "Climate-Spring":   ("spring-bloom",       "spring-bloom"),
    "Climate-Rainy":    ("rainy-day",          "rainy-day"),
    "Climate-Humid":    ("humid-weather",      "humid-weather"),
    "Climate-Tropical": ("tropical-climate",   "tropical-climate"),
    "Climate-Desert":   ("desert-climate",     "desert-climate"),
    "Climate-Hot":      ("hot-weather",        "hot-weather"),
    "Climate-Cold":     ("cold-weather",       "cold-weather"),
    "Climate-Dry":      ("dry-weather",        "dry-weather"),
    "Arabic Perfumes":  ("arabic",             "arabic"),
    "Luxury Perfumes":  ("luxury",             "luxury"),
    "Crazy Deals":      ("crazy-deals",        "sale"),
}
# "Needs Review" is deliberately absent: it is the spreadsheet's leftovers pile,
# not a collection in the store.

MANAGED_TAGS = {tag for _, tag in SHEETS.values()}
HEADER_ROWS = 4          # title, count line, blank, column headers
HANDLE_COL = 7           # "Handle (Shopify URL slug)"


# ── plumbing ─────────────────────────────────────────────────────────────
def load_env():
    env = {}
    path = REPO / ".env"
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                env[k.strip()] = v.strip()
    dom = env.get("SHOPIFY_STORE_DOMAIN") or os.environ.get("SHOPIFY_STORE_DOMAIN", "")
    tok = env.get("SHOPIFY_ADMIN_ACCESS_TOKEN") or os.environ.get("SHOPIFY_ADMIN_ACCESS_TOKEN", "")
    ver = env.get("SHOPIFY_API_VERSION") or os.environ.get("SHOPIFY_API_VERSION", "2024-10")
    dom = dom.replace("https://", "").replace("http://", "").strip("/")
    if not dom or not tok:
        sys.exit("Missing SHOPIFY_STORE_DOMAIN / SHOPIFY_ADMIN_ACCESS_TOKEN (.env or environment)")
    return dom, tok, ver


try:
    import certifi
    SSL_CTX = ssl.create_default_context(cafile=certifi.where())
except ImportError:
    SSL_CTX = ssl.create_default_context()


def request(method, url, token, payload=None):
    data = json.dumps(payload).encode() if payload is not None else None
    req = urllib.request.Request(url, data=data, method=method, headers={
        "X-Shopify-Access-Token": token,
        "Accept": "application/json",
        "Content-Type": "application/json",
    })
    for attempt in range(6):
        try:
            resp = urllib.request.urlopen(req, timeout=60, context=SSL_CTX)
            return json.load(resp), resp.headers
        except urllib.error.HTTPError as e:
            if e.code == 429:                       # leaky bucket is full
                time.sleep(float(e.headers.get("Retry-After", 2)))
                continue
            raise
    raise RuntimeError(f"giving up after retries: {url}")


def fetch_products(dom, tok, ver):
    """Every product with the fields we need, following Link pagination."""
    out, url = [], (f"https://{dom}/admin/api/{ver}/products.json"
                    f"?limit=250&fields=id,handle,title,tags,status")
    while url:
        body, headers = request("GET", url, tok)
        out.extend(body["products"])
        m = re.search(r'<([^>]+)>;\s*rel="next"', headers.get("Link", ""))
        url = m.group(1) if m else None
    return out


def read_sheets():
    """handle -> set of tags the spreadsheet says it should carry, plus per-sheet stats."""
    try:
        import openpyxl
    except ImportError:
        sys.exit("pip install openpyxl")
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    want = defaultdict(set)
    stats = {}
    for sheet, (coll, tag) in SHEETS.items():
        if sheet not in wb.sheetnames:
            sys.exit(f"sheet missing from workbook: {sheet}")
        ws = wb[sheet]
        stated, rows, handles = None, 0, set()
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i == 1 and r and r[0]:                      # "N products matched ..."
                m = re.match(r"\s*(\d+)", str(r[0]))
                stated = int(m.group(1)) if m else None
            if i < HEADER_ROWS or not r or not r[0]:
                continue
            rows += 1
            if len(r) > HANDLE_COL and r[HANDLE_COL]:
                h = str(r[HANDLE_COL]).strip()
                handles.add(h)
                want[h].add(tag)
        stats[sheet] = {"collection": coll, "tag": tag, "stated": stated,
                        "rows": rows, "handles": handles}
    return want, stats


# ── main ─────────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="write the changes")
    ap.add_argument("--collection", action="append", default=None,
                    help="limit to these collection handles (repeatable)")
    ap.add_argument("--report", help="write the dry-run report to this file")
    args = ap.parse_args()

    dom, tok, ver = load_env()
    sheets = SHEETS
    if args.collection:
        keep = set(args.collection)
        sheets = {s: v for s, v in SHEETS.items() if v[0] in keep}
        if not sheets:
            sys.exit(f"no sheet maps to: {', '.join(keep)}")
    scoped_tags = {tag for _, tag in sheets.values()}

    L = []
    def say(s=""):
        print(s)
        L.append(s)

    say(f"store   : {dom}")
    say(f"source  : {XLSX.name}")
    say(f"scope   : {len(sheets)} collections, {len(scoped_tags)} managed tags")
    say()

    want, stats = read_sheets()
    products = fetch_products(dom, tok, ver)
    by_handle = {p["handle"]: p for p in products}
    say(f"store has {len(products)} products; spreadsheet names "
        f"{len({h for st in stats.values() for h in st['handles']})} handles")
    say()

    # ── integrity: does each sheet agree with its own stated count? ──────
    say("## Sheet integrity")
    say(f"{'sheet':<18}{'states':>8}{'rows':>7}{'unique':>8}{'in store':>10}  flag")
    suspect = []
    for sheet, st in stats.items():
        if sheet not in sheets:
            continue
        present = len(st["handles"] & by_handle.keys())
        flags = []
        if st["stated"] is not None and st["stated"] != st["rows"]:
            flags.append("COUNT MISMATCH")
            suspect.append((sheet, st))
        if len(st["handles"]) != st["rows"]:
            flags.append(f"{st['rows']-len(st['handles'])} dupes")
        say(f"{sheet:<18}{st['stated'] if st['stated'] is not None else '?':>8}"
            f"{st['rows']:>7}{len(st['handles']):>8}{present:>10}  {', '.join(flags)}")
    say()

    # ── the actual diff ──────────────────────────────────────────────────
    adds = defaultdict(list)
    removes = defaultdict(list)
    changed = {}
    for p in products:
        cur = {t.strip() for t in (p["tags"] or "").split(",") if t.strip()}
        desired = want.get(p["handle"], set()) & scoped_tags
        cur_scoped = cur & scoped_tags
        add, rm = desired - cur_scoped, cur_scoped - desired
        if not add and not rm:
            continue
        for t in add:
            adds[t].append(p["handle"])
        for t in rm:
            removes[t].append(p["handle"])
        # everything outside the managed set is carried through untouched
        changed[p["id"]] = {"handle": p["handle"], "title": p["title"],
                            "tags": sorted((cur - scoped_tags) | desired),
                            "add": sorted(add), "remove": sorted(rm)}

    say("## Membership change per collection")
    say(f"{'collection':<20}{'tag':<24}{'now':>6}{'after':>7}{'+':>7}{'-':>7}")
    for sheet, (coll, tag) in sheets.items():
        now = sum(1 for p in products
                  if tag in {t.strip() for t in (p["tags"] or "").split(",")})
        after = now + len(adds[tag]) - len(removes[tag])
        say(f"{coll:<20}{tag:<24}{now:>6}{after:>7}{len(adds[tag]):>+7}{-len(removes[tag]):>+7}")
    say()

    missing = {h for st in stats.values() if st["collection"] in {c for c, _ in sheets.values()}
               for h in st["handles"]} - by_handle.keys()
    say(f"products touched      : {len(changed)}")
    say(f"tag additions         : {sum(len(v) for v in adds.values())}")
    say(f"tag removals          : {sum(len(v) for v in removes.values())}")
    say(f"handles not in store  : {len(missing)}  (silently skipped)")
    say()

    if suspect:
        say("## WARNING — sheets that contradict their own header")
        for sheet, st in suspect:
            say(f"  {sheet}: header says {st['stated']}, sheet lists {st['rows']} rows"
                f"  ({st['rows']/max(st['stated'],1):.0f}x)")
        say()

    if args.report:
        Path(args.report).write_text("\n".join(L) + "\n", encoding="utf-8")
        print(f"[report written to {args.report}]")

    if not args.apply:
        print("[dry run — nothing written. Re-run with --apply to write tags.]")
        return

    print(f"\nAbout to write `tags` on {len(changed)} products on {dom}.")
    if input("Type the store domain to confirm: ").strip() != dom:
        sys.exit("confirmation did not match — nothing written")

    for n, (pid, c) in enumerate(changed.items(), 1):
        request("PUT", f"https://{dom}/admin/api/{ver}/products/{pid}.json", tok,
                {"product": {"id": pid, "tags": ", ".join(c["tags"])}})
        print(f"  [{n}/{len(changed)}] {c['handle']}  +{c['add']} -{c['remove']}")
        time.sleep(0.55)                      # 2 calls/sec limit, with headroom
    print("done")


if __name__ == "__main__":
    main()
