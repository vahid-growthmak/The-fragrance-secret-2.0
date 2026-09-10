#!/usr/bin/env python3
"""Fill the SEO fields and image alt text from the final catalogue sheet.

The metafield import covers everything the PDP body reads. This covers the
three sheet columns that land somewhere other than a custom metafield:

    sheet column            destination
    Meta Title              product.seo.title       (the <title> and SERP link)
    Meta Description        product.seo.description (the SERP snippet)
    Image Alt Text (SEO)    alt text on the product's first image

and one that is a metafield but belongs with them:

    FAQ Schema (JSON-LD)    custom.faq_schema

Deliberately not imported, with reasons:

    Product Schema (JSON-LD)
        snippets/product-schema.liquid builds Product structured data from live
        product data on purpose, and says why: the stored copy carries a null
        price, a hardcoded availability and a fixed image list. Checked against
        the sheet and the concern is real — 1,015 of the 1,308 stored schemas
        have a null price, and all of them declare AED while this store is set
        to USD. Emitting that would be invalid structured data. The FAQ schema
        has no such problem, which is why the snippet emits that one as
        supplied and only rebuilds the Product one.

    Canonical URL
        every row points at thefragrancesecrets.com, a different store from
        this one. layout/theme.liquid already emits {{ canonical_url }}, which
        resolves correctly per store. Importing these would hardcode canonicals
        to another domain.

    Focus Keyword, Secondary Keywords
        nothing in the theme reads them and Shopify has no field for them. They
        are working notes for whoever writes the copy, not product data.

Only values that differ are written, so a second run is a no-op.

    python3 scripts/import-product-seo.py            # dry run
    python3 scripts/import-product-seo.py --limit 5  # dry run, first 5
    python3 scripts/import-product-seo.py --apply    # write
"""

import argparse
import json
import os
import re
import ssl
import sys
import time
import urllib.error
import urllib.request
from collections import Counter
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
XLSX = REPO / "TFS Final Catalog Sheet - 1_308 ProductsSEO_AEO_GEO Optimized.xlsx"

try:
    import certifi
    SSL_CTX = ssl.create_default_context(cafile=certifi.where())
except ImportError:
    SSL_CTX = ssl.create_default_context()


def load_env():
    env = {}
    p = REPO / ".env"
    for line in (p.read_text(encoding="utf-8").splitlines() if p.exists() else []):
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip()
    dom = (env.get("SHOPIFY_STORE_DOMAIN") or os.environ.get("SHOPIFY_STORE_DOMAIN", ""))
    tok = (env.get("SHOPIFY_ADMIN_ACCESS_TOKEN") or os.environ.get("SHOPIFY_ADMIN_ACCESS_TOKEN", ""))
    ver = (env.get("SHOPIFY_API_VERSION") or os.environ.get("SHOPIFY_API_VERSION", "2024-10"))
    dom = dom.replace("https://", "").replace("http://", "").strip("/")
    if not dom or not tok:
        sys.exit("Missing SHOPIFY_STORE_DOMAIN / SHOPIFY_ADMIN_ACCESS_TOKEN")
    return dom, tok, ver


def http(method, url, token, payload=None):
    data = json.dumps(payload).encode() if payload is not None else None
    req = urllib.request.Request(url, data=data, method=method, headers={
        "X-Shopify-Access-Token": token, "Accept": "application/json",
        "Content-Type": "application/json"})
    for attempt in range(8):
        try:
            r = urllib.request.urlopen(req, timeout=120, context=SSL_CTX)
            body = r.read()
            return json.loads(body) if body.strip() else {}
        except urllib.error.HTTPError as e:
            if e.code == 429:
                time.sleep(float(e.headers.get("Retry-After", 2)))
                continue
            if 500 <= e.code < 600:
                time.sleep(2 ** attempt)
                continue
            raise
        except (urllib.error.URLError, TimeoutError, ConnectionError, ssl.SSLError):
            if attempt == 7:
                raise
            time.sleep(2 ** attempt)
    raise RuntimeError(f"retries exhausted: {url}")


def gql(dom, tok, ver, query, variables=None):
    d = http("POST", f"https://{dom}/admin/api/{ver}/graphql.json", tok,
             {"query": query, "variables": variables or {}})
    if d.get("errors"):
        raise SystemExit(json.dumps(d["errors"])[:500])
    return d["data"]


def norm(s):
    s = re.sub(r"[–—]", "-", str(s or "").lower())
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def read_sheet():
    try:
        import openpyxl
    except ImportError:
        sys.exit("pip install openpyxl")
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    hdr = [str(h).strip() if h else "" for h in rows[0]]
    idx = {h: i for i, h in enumerate(hdr)}
    need = ["title", "Meta Title", "Meta Description", "Image Alt Text (SEO)",
            "FAQ Schema (JSON-LD)"]
    missing = [c for c in need if c not in idx]
    if missing:
        sys.exit(f"sheet is missing expected columns: {missing}")
    out = {}
    for r in rows[1:]:
        if not r or not r[0] or not str(r[0]).strip():
            continue
        def g(c):
            i = idx[c]
            return "" if i >= len(r) or r[i] is None else str(r[i]).strip()
        out[norm(g("title"))] = {
            "seo_title": g("Meta Title"),
            "seo_desc": g("Meta Description"),
            "alt": g("Image Alt Text (SEO)"),
            "faq_schema": g("FAQ Schema (JSON-LD)"),
        }
    return out


PRODUCTS_Q = """query($c:String){ products(first:150, after:$c){
  pageInfo{hasNextPage endCursor}
  nodes{ id title seo{title description}
    media(first:1){ nodes{ id alt } }
    metafields(first:45, namespace:"custom"){ nodes{ key value } } } } }"""

UPDATE_SEO = """mutation($id:ID!,$seo:SEOInput!){
  productUpdate(input:{id:$id, seo:$seo}){ userErrors{ field message } } }"""

UPDATE_ALT = """mutation($pid:ID!,$media:[UpdateMediaInput!]!){
  productUpdateMedia(productId:$pid, media:$media){ mediaUserErrors{ field message } } }"""

SET_MF = """mutation($m:[MetafieldsSetInput!]!){
  metafieldsSet(metafields:$m){ userErrors{ field message } } }"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--limit", type=int)
    args = ap.parse_args()
    dom, tok, ver = load_env()

    sheet = read_sheet()
    products, cur = [], None
    while True:
        page = gql(dom, tok, ver, PRODUCTS_Q, {"c": cur})["products"]
        products += page["nodes"]
        if not page["pageInfo"]["hasNextPage"]:
            break
        cur = page["pageInfo"]["endCursor"]
    print(f"store : {dom}  ({len(products)} products)")
    print(f"sheet : {XLSX.name}  ({len(sheet)} rows)\n")

    plan, unmatched = [], 0
    counts = Counter()
    for p in products:
        row = sheet.get(norm(p["title"]))
        if not row:
            unmatched += 1
            continue
        job = {"id": p["id"], "title": p["title"]}
        # SEOInput is not a patch: a field left out is written as null. Sending
        # {title} alone on a product whose description already matched wiped
        # that description — it did exactly that to 177 products before this
        # was caught. Always send both fields, sourcing each from the sheet and
        # falling back to whatever the product already has.
        want_title = row["seo_title"][:70] if row["seo_title"] else (p["seo"]["title"] or "")
        want_desc = row["seo_desc"][:320] if row["seo_desc"] else (p["seo"]["description"] or "")
        if want_title != (p["seo"]["title"] or "") or want_desc != (p["seo"]["description"] or ""):
            job["seo"] = {"title": want_title, "description": want_desc}
            if want_title != (p["seo"]["title"] or ""):
                counts["seo_title"] += 1
            if want_desc != (p["seo"]["description"] or ""):
                counts["seo_description"] += 1
        media = p["media"]["nodes"]
        if row["alt"] and media and (media[0]["alt"] or "") != row["alt"]:
            job["alt"] = {"id": media[0]["id"], "alt": row["alt"][:512]}
            counts["image alt"] += 1
        have_mf = {m["key"]: (m["value"] or "") for m in p["metafields"]["nodes"]}
        if row["faq_schema"] and have_mf.get("faq_schema", "") != row["faq_schema"]:
            job["faq_schema"] = row["faq_schema"]
            counts["faq_schema"] += 1
        if len(job) > 2:
            plan.append(job)
    if args.limit:
        plan = plan[: args.limit]

    print(f"products needing changes : {len(plan)}")
    print(f"already correct          : {len(products) - unmatched - len(plan)}")
    print(f"not in the sheet         : {unmatched}\n")
    for k, c in counts.most_common():
        print(f"  {k:<16}{c:>6}")

    if not args.apply:
        print("\n[dry run — nothing written. Re-run with --apply.]")
        return

    print(f"\nAbout to write SEO fields, alt text and FAQ schema on {len(plan)} products.")
    if input("Type the store domain to confirm: ").strip() != dom:
        sys.exit("confirmation did not match — nothing written")

    ok = failed = 0
    errors = []
    log = REPO / "scripts" / "seo-import-run.json"
    for n, job in enumerate(plan, 1):
        try:
            if job.get("seo"):
                res = gql(dom, tok, ver, UPDATE_SEO, {"id": job["id"], "seo": job["seo"]})
                errs = res["productUpdate"]["userErrors"]
                if errs:
                    raise RuntimeError(errs[:2])
            if job.get("alt"):
                res = gql(dom, tok, ver, UPDATE_ALT, {
                    "pid": job["id"],
                    "media": [{"id": job["alt"]["id"], "alt": job["alt"]["alt"]}]})
                errs = res["productUpdateMedia"]["mediaUserErrors"]
                if errs:
                    raise RuntimeError(errs[:2])
            if job.get("faq_schema"):
                res = gql(dom, tok, ver, SET_MF, {"m": [{
                    "ownerId": job["id"], "namespace": "custom", "key": "faq_schema",
                    "type": "multi_line_text_field", "value": job["faq_schema"][:65000]}]})
                errs = res["metafieldsSet"]["userErrors"]
                if errs:
                    raise RuntimeError(errs[:2])
            ok += 1
        except Exception as e:
            failed += 1
            errors.append({"title": job["title"][:70], "error": str(e)[:200]})
        if n % 50 == 0 or n == len(plan):
            print(f"  [{n}/{len(plan)}] {ok} ok, {failed} failed", flush=True)
            log.write_text(json.dumps({"ok": ok, "failed": failed, "errors": errors[:40]},
                                      indent=1), encoding="utf-8")
        time.sleep(0.4)
    log.write_text(json.dumps({"ok": ok, "failed": failed, "errors": errors[:40]},
                              indent=1), encoding="utf-8")
    print(f"\n{ok} products updated, {failed} failed -> {log}")


if __name__ == "__main__":
    main()
