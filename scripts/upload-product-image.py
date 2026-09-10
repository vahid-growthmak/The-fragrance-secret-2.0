#!/usr/bin/env python3
"""Attach a local image file to a product as its featured image.

sync-paris-images.py ingests images by URL, which only works for files already
hosted somewhere public. A file on this machine has to go up through Shopify's
staged upload flow instead:

    stagedUploadsCreate  ->  POST the bytes to the returned bucket URL
                         ->  productCreateMedia with the returned resourceUrl

The new image is placed first, so it becomes the featured image on cards, the
PDP hero and the Google feed. Existing images are kept and shuffled down
rather than replaced — pass --replace to drop them instead.

Alt text comes from the sheet's Image Alt Text (SEO) column, matching what the
SEO import wrote, so this does not undo that work.

    python3 scripts/upload-product-image.py --set handle=path/to.jpg
    python3 scripts/upload-product-image.py --set handle=path/to.jpg --apply
"""

import argparse
import json
import mimetypes
import os
import re
import ssl
import sys
import time
import urllib.error
import urllib.request
import uuid
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
XLSX = REPO / "TFS Final Catalog Sheet - 1_308 ProductsSEO_AEO_GEO Optimized.xlsx"

try:
    import certifi
    SSL_CTX = ssl.create_default_context(cafile=certifi.where())
except ImportError:
    SSL_CTX = ssl.create_default_context()


def load_env():
    env = {}
    p = REPO / ".env"
    for line in (p.read_text(encoding="utf-8").splitlines() if p.exists() else []):
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip()
    dom = (env.get("SHOPIFY_STORE_DOMAIN") or os.environ.get("SHOPIFY_STORE_DOMAIN", ""))
    tok = (env.get("SHOPIFY_ADMIN_ACCESS_TOKEN") or os.environ.get("SHOPIFY_ADMIN_ACCESS_TOKEN", ""))
    ver = (env.get("SHOPIFY_API_VERSION") or os.environ.get("SHOPIFY_API_VERSION", "2024-10"))
    dom = dom.replace("https://", "").replace("http://", "").strip("/")
    if not dom or not tok:
        sys.exit("Missing SHOPIFY_STORE_DOMAIN / SHOPIFY_ADMIN_ACCESS_TOKEN")
    return dom, tok, ver


def gql(dom, tok, ver, query, variables=None):
    body = json.dumps({"query": query, "variables": variables or {}}).encode()
    req = urllib.request.Request(f"https://{dom}/admin/api/{ver}/graphql.json", data=body,
                                 method="POST", headers={
                                     "X-Shopify-Access-Token": tok,
                                     "Content-Type": "application/json"})
    for attempt in range(6):
        try:
            d = json.loads(urllib.request.urlopen(req, timeout=120, context=SSL_CTX).read())
            break
        except urllib.error.HTTPError as e:
            if e.code == 429 or (500 <= e.code < 600 and attempt < 5):
                time.sleep(2 ** attempt)
                continue
            raise
        except (urllib.error.URLError, TimeoutError, ConnectionError, ssl.SSLError):
            if attempt == 5:
                raise
            time.sleep(2 ** attempt)
    else:
        raise RuntimeError("retries exhausted")
    if d.get("errors"):
        raise SystemExit(json.dumps(d["errors"])[:600])
    return d["data"]


def post_multipart(url, fields, filename, blob, content_type):
    """Shopify's bucket requires the form fields in the given order, file last."""
    boundary = f"----tfs{uuid.uuid4().hex}"
    parts = []
    for k, v in fields:
        parts.append(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n"
                     .encode())
    parts.append(
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; "
        f"filename=\"{filename}\"\r\nContent-Type: {content_type}\r\n\r\n".encode())
    parts.append(blob)
    parts.append(f"\r\n--{boundary}--\r\n".encode())
    body = b"".join(parts)
    req = urllib.request.Request(url, data=body, method="POST", headers={
        "Content-Type": f"multipart/form-data; boundary={boundary}",
        "Content-Length": str(len(body))})
    with urllib.request.urlopen(req, timeout=300, context=SSL_CTX) as r:
        return r.status, r.read()[:300]


def sheet_alt():
    try:
        import openpyxl
    except ImportError:
        return {}
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    hdr = [str(h).strip() if h else "" for h in rows[0]]
    idx = {h: i for i, h in enumerate(hdr)}
    out = {}
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        def g(c):
            i = idx.get(c)
            return "" if i is None or i >= len(r) or r[i] is None else str(r[i]).strip()
        h = g("Canonical URL").rstrip("/").split("/")[-1]
        if h:
            out[h] = g("Image Alt Text (SEO)")
    return out


PRODUCT_Q = """query($h:String!){ productByHandle(handle:$h){ id title
  media(first:50){ nodes{ id status mediaContentType
    ... on MediaImage{ image{ url width height } } } } } }"""

STAGE = """mutation($input:[StagedUploadInput!]!){ stagedUploadsCreate(input:$input){
  stagedTargets{ url resourceUrl parameters{ name value } } userErrors{ field message } } }"""

CREATE = """mutation($pid:ID!,$m:[CreateMediaInput!]!){
  productCreateMedia(productId:$pid, media:$m){
    media{ ... on MediaImage{ id status } } mediaUserErrors{ field message } } }"""

REORDER = """mutation($id:ID!,$moves:[MoveInput!]!){
  productReorderMedia(id:$id, moves:$moves){ mediaUserErrors{ field message } } }"""

DELETE = """mutation($pid:ID!,$ids:[ID!]!){
  productDeleteMedia(productId:$pid, mediaIds:$ids){ mediaUserErrors{ field message } } }"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--set", action="append", required=True, metavar="HANDLE=FILE",
                    help="product handle and the local image to make its featured image")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--replace", action="store_true",
                    help="remove the product's other images instead of keeping them")
    args = ap.parse_args()
    dom, tok, ver = load_env()
    alts = sheet_alt()

    jobs = []
    for spec in args.set:
        if "=" not in spec:
            sys.exit(f"--set needs HANDLE=FILE, got {spec!r}")
        handle, path = spec.split("=", 1)
        f = Path(path)
        if not f.exists():
            sys.exit(f"no such file: {f}")
        p = gql(dom, tok, ver, PRODUCT_Q, {"h": handle})["productByHandle"]
        if not p:
            sys.exit(f"no product with handle {handle!r}")
        imgs = [m for m in p["media"]["nodes"] if m["mediaContentType"] == "IMAGE"]
        jobs.append({"handle": handle, "file": f, "product": p, "existing": imgs})
        print(f"{p['title'][:66]}")
        print(f"   handle    {handle}")
        print(f"   new image {f.name}  ({f.stat().st_size/1024:.0f} KB)")
        print(f"   currently {len(imgs)} image(s); new one becomes #1, "
              f"{'others removed' if args.replace else 'others shift down'}")
        print(f"   alt       {alts.get(handle,'(none in sheet)')[:64]!r}\n")

    if not args.apply:
        print("[dry run — nothing uploaded. Re-run with --apply.]")
        return

    print(f"About to upload {len(jobs)} image(s) to {dom}.")
    if input("Type the store domain to confirm: ").strip() != dom:
        sys.exit("confirmation did not match — nothing uploaded")

    for j in jobs:
        f, p = j["file"], j["product"]
        blob = f.read_bytes()
        ctype = mimetypes.guess_type(f.name)[0] or "image/jpeg"
        print(f"\n{p['title'][:60]}")

        tgt = gql(dom, tok, ver, STAGE, {"input": [{
            "filename": f.name, "mimeType": ctype, "resource": "IMAGE",
            "httpMethod": "POST", "fileSize": str(len(blob))}]})["stagedUploadsCreate"]
        if tgt.get("userErrors"):
            print(f"   STAGE FAILED {tgt['userErrors'][:2]}")
            continue
        t = tgt["stagedTargets"][0]
        status, resp = post_multipart(t["url"], [(x["name"], x["value"]) for x in t["parameters"]],
                                      f.name, blob, ctype)
        if status not in (200, 201, 204):
            print(f"   UPLOAD FAILED http {status} {resp!r}")
            continue
        print(f"   uploaded ({len(blob)/1024:.0f} KB)")

        alt = alts.get(j["handle"], "")
        res = gql(dom, tok, ver, CREATE, {"pid": p["id"], "m": [{
            "originalSource": t["resourceUrl"], "mediaContentType": "IMAGE",
            "alt": alt[:512]}]})["productCreateMedia"]
        if res.get("mediaUserErrors"):
            print(f"   ATTACH FAILED {res['mediaUserErrors'][:2]}")
            continue
        new_id = res["media"][0]["id"]

        for _ in range(40):
            nodes = gql(dom, tok, ver, PRODUCT_Q, {"h": j["handle"]})["productByHandle"]["media"]["nodes"]
            me = next((n for n in nodes if n["id"] == new_id), None)
            if me and me["status"] == "READY":
                break
            if me and me["status"] == "FAILED":
                print("   PROCESSING FAILED")
                me = None
                break
            time.sleep(3)
        if not me:
            continue
        print("   processed")

        if args.replace:
            old = [m["id"] for m in j["existing"]]
            if old:
                gql(dom, tok, ver, DELETE, {"pid": p["id"], "ids": old})
                print(f"   removed {len(old)} previous image(s)")
        else:
            gql(dom, tok, ver, REORDER, {"id": p["id"],
                                         "moves": [{"id": new_id, "newPosition": "0"}]})
            print("   moved to position 1 (featured)")
        time.sleep(0.4)

    print("\ndone")


if __name__ == "__main__":
    main()
